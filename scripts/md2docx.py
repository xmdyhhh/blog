#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Convert Markdown experiment report to Word document."""

import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def set_cell_shading(cell, color):
    """Set cell background color."""
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), color)
    cell._tc.get_or_add_tcPr().append(shading)

def add_formatted_text(paragraph, text):
    """Add text with inline formatting (bold, italic, code)."""
    # Pattern for **bold**, *italic*, `code`
    pattern = r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)'
    parts = re.split(pattern, text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        elif part.startswith('`') and part.endswith('`'):
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Consolas'
            run.font.color.rgb = RGBColor(0xC7, 0x25, 0x4E)
            run.font.size = Pt(10)
        else:
            paragraph.add_run(part)

def parse_table(lines, start_idx):
    """Parse a markdown table starting at start_idx. Returns (table_rows, end_idx)."""
    rows = []
    i = start_idx
    while i < len(lines) and lines[i].strip().startswith('|'):
        line = lines[i].strip()
        # Skip separator line
        if re.match(r'^\|[\s\-:|]+\|$', line):
            i += 1
            continue
        cells = [c.strip() for c in line.split('|')[1:-1]]
        rows.append(cells)
        i += 1
    return rows, i

def md_to_docx(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()
    
    lines = content.split('\n')
    doc = Document()
    
    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = '宋体'
    font.size = Pt(12)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    
    # Set margins
    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(3.17)
        section.right_margin = Cm(3.17)
    
    i = 0
    in_code_block = False
    code_lines = []
    
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        
        # Code block
        if stripped.startswith('```'):
            if in_code_block:
                # End code block
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.5)
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(6)
                run = p.add_run('\n'.join(code_lines))
                run.font.name = 'Consolas'
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
                # Add shading
                pPr = p._p.get_or_add_pPr()
                shd = OxmlElement('w:shd')
                shd.set(qn('w:fill'), 'F5F5F5')
                pPr.append(shd)
                code_lines = []
                in_code_block = False
            else:
                in_code_block = True
            i += 1
            continue
        
        if in_code_block:
            code_lines.append(line)
            i += 1
            continue
        
        # Empty line
        if not stripped:
            i += 1
            continue
        
        # Horizontal rule
        if stripped == '---' or stripped == '***':
            p = doc.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            pBdr = OxmlElement('w:pBdr')
            bottom = OxmlElement('w:bottom')
            bottom.set(qn('w:val'), 'single')
            bottom.set(qn('w:sz'), '6')
            bottom.set(qn('w:space'), '1')
            bottom.set(qn('w:color'), '999999')
            pBdr.append(bottom)
            pPr.append(pBdr)
            i += 1
            continue
        
        # Headings
        if stripped.startswith('#'):
            match = re.match(r'^(#{1,6})\s+(.*)$', stripped)
            if match:
                level = len(match.group(1))
                text = match.group(2)
                if level == 1:
                    p = doc.add_heading('', level=0)
                    run = p.add_run(text)
                    run.font.size = Pt(22)
                    run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                elif level == 2:
                    p = doc.add_heading('', level=1)
                    run = p.add_run(text)
                    run.font.size = Pt(16)
                    run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
                elif level == 3:
                    p = doc.add_heading('', level=2)
                    run = p.add_run(text)
                    run.font.size = Pt(14)
                    run.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
                else:
                    p = doc.add_heading('', level=3)
                    run = p.add_run(text)
                    run.font.size = Pt(12)
                # Set East Asian font for headings
                for r in p.runs:
                    r.font.name = '黑体'
                    r._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
                i += 1
                continue
        
        # Table
        if stripped.startswith('|') and '|' in stripped[1:]:
            rows, end_idx = parse_table(lines, i)
            if rows and len(rows) > 0:
                num_cols = max(len(r) for r in rows)
                table = doc.add_table(rows=len(rows), cols=num_cols)
                table.style = 'Table Grid'
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                for row_idx, row_cells in enumerate(rows):
                    for col_idx, cell_text in enumerate(row_cells):
                        if col_idx < num_cols:
                            cell = table.cell(row_idx, col_idx)
                            cell.text = ''
                            p = cell.paragraphs[0]
                            add_formatted_text(p, cell_text)
                            for run in p.runs:
                                run.font.size = Pt(10)
                                run.font.name = '宋体'
                                run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
                            if row_idx == 0:
                                set_cell_shading(cell, '1E3A8A')
                                for run in p.runs:
                                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                                    run.bold = True
                i = end_idx
                doc.add_paragraph()  # spacing after table
                continue
        
        # Unordered list
        if re.match(r'^[\-\*\+]\s+', stripped):
            text = re.sub(r'^[\-\*\+]\s+', '', stripped)
            p = doc.add_paragraph(style='List Bullet')
            add_formatted_text(p, text)
            for run in p.runs:
                run.font.size = Pt(12)
                run.font.name = '宋体'
                run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
            i += 1
            continue
        
        # Ordered list
        if re.match(r'^\d+\.\s+', stripped):
            text = re.sub(r'^\d+\.\s+', '', stripped)
            p = doc.add_paragraph(style='List Number')
            add_formatted_text(p, text)
            for run in p.runs:
                run.font.size = Pt(12)
                run.font.name = '宋体'
                run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
            i += 1
            continue
        
        # Normal paragraph
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0.74)
        p.paragraph_format.line_spacing = 1.5
        add_formatted_text(p, stripped)
        for run in p.runs:
            run.font.size = Pt(12)
            run.font.name = '宋体'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        
        i += 1
    
    doc.save(docx_path)
    print(f"Word document saved to: {docx_path}")

if __name__ == '__main__':
    md_path = r'D:\Desktop\开源软件作业\oss-blog\docs\实验报告.md'
    docx_path = r'D:\Desktop\开源软件作业\oss-blog\docs\实验报告.docx'
    md_to_docx(md_path, docx_path)
